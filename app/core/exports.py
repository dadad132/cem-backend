"""Excel and PDF files for the Export buttons (All Tasks, Board View, projects).

An export is a title, a few lines saying what it covers, some summary figures
and one or more tables. Each page only gathers its rows; the look matches the
ticket report (indigo header rows, auto-sized Excel columns).
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any, List, Optional, Sequence, Tuple
from xml.sax.saxutils import escape

from fastapi.responses import Response

HEADER_COLOR = '6366F1'  # as the ticket report
FORMATS = ('pdf', 'excel')


@dataclass
class ExportTable:
    title: str                      # Excel sheet name / PDF section heading
    columns: List[str]
    rows: List[Sequence[Any]]
    # Relative PDF column widths (wide text columns get more); equal if None
    widths: Optional[List[float]] = None


@dataclass
class Export:
    title: str
    filename: str                   # without extension
    lines: List[str] = field(default_factory=list)              # under the title
    summary: List[Tuple[str, Any]] = field(default_factory=list)  # (figure, value)
    tables: List[ExportTable] = field(default_factory=list)


def _text(value) -> str:
    if value is None:
        return ''
    if isinstance(value, datetime):
        return value.strftime('%Y-%m-%d %H:%M')
    if isinstance(value, date):
        return value.strftime('%Y-%m-%d')
    return str(value)


def to_excel(export: Export) -> bytes:
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color=HEADER_COLOR, end_color=HEADER_COLOR, fill_type='solid')
    thin = Side(style='thin', color='D1D5DB')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def clean(value):
        # Control characters (e.g. from e-mailed text) make Excel refuse the file
        return ILLEGAL_CHARACTERS_RE.sub('', value) if isinstance(value, str) else value

    def style_header(ws, row, count):
        for col in range(1, count + 1):
            cell = ws.cell(row=row, column=col)
            cell.font, cell.fill, cell.border = header_font, header_fill, border
            cell.alignment = Alignment(horizontal='center', vertical='center')

    def fit_columns(ws):
        for cells in ws.columns:
            longest = max((len(_text(c.value)) for c in cells if c.value is not None), default=0)
            ws.column_dimensions[cells[0].column_letter].width = min(max(longest + 2, 8), 60)

    wb = Workbook()
    ws = wb.active
    ws.title = 'Summary'
    ws.append([clean(export.title)])
    ws.cell(row=1, column=1).font = Font(bold=True, size=14)
    for line in export.lines:
        ws.append([clean(line)])
    if export.summary:
        ws.append([])
        ws.append(['Figure', 'Value'])
        style_header(ws, ws.max_row, 2)
        for label, value in export.summary:
            ws.append([clean(label), clean(value)])
    fit_columns(ws)

    used = {'Summary'}
    for table in export.tables:
        name = re.sub(r'[\[\]:*?/\\]', '', table.title)[:31] or 'Sheet'
        base, n = name, 2
        while name in used:
            name = f'{base[:28]} {n}'
            n += 1
        used.add(name)
        sheet = wb.create_sheet(name)
        sheet.append(table.columns)
        style_header(sheet, 1, len(table.columns))
        for row in table.rows:
            sheet.append([clean(v) for v in row])
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, datetime):
                    cell.number_format = 'yyyy-mm-dd hh:mm'
                elif isinstance(cell.value, date):
                    cell.number_format = 'yyyy-mm-dd'
        sheet.freeze_panes = 'A2'
        if table.rows:
            sheet.auto_filter.ref = sheet.dimensions
        fit_columns(sheet)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def to_pdf(export: Export) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    base = getSampleStyleSheet()['Normal']
    title_style = ParagraphStyle('t', parent=base, fontName='Helvetica-Bold', fontSize=16, leading=20,
                                 textColor=colors.HexColor('#1e293b'))
    line_style = ParagraphStyle('l', parent=base, fontSize=9, leading=12, textColor=colors.HexColor('#64748b'))
    heading_style = ParagraphStyle('h', parent=base, fontName='Helvetica-Bold', fontSize=11, leading=14,
                                   textColor=colors.HexColor('#1e293b'), spaceBefore=8, spaceAfter=4)
    head_cell = ParagraphStyle('hc', parent=base, fontName='Helvetica-Bold', fontSize=8, leading=10,
                               textColor=colors.white)
    body_cell = ParagraphStyle('bc', parent=base, fontSize=8, leading=10, textColor=colors.HexColor('#0f172a'))

    def cell(value, style):
        # Paragraphs wrap long text; escaped so "<" or "&" print as typed
        return Paragraph(escape(_text(value)).replace('\n', '<br/>'), style)

    page = landscape(A4)
    width = page[0] - 24 * mm

    def table(columns, rows, weights=None):
        weights = weights or [1] * len(columns)
        col_widths = [width * w / sum(weights) for w in weights]
        data = [[cell(c, head_cell) for c in columns]] + [[cell(v, body_cell) for v in r] for r in rows]
        t = Table(data, colWidths=col_widths, repeatRows=1)
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#' + HEADER_COLOR)),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')]),
            ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#e2e8f0')),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('TOPPADDING', (0, 0), (-1, -1), 3), ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ]))
        return t

    story = [Paragraph(escape(export.title), title_style), Spacer(1, 2 * mm)]
    story += [Paragraph(escape(line), line_style) for line in export.lines]
    if export.summary:
        story += [Spacer(1, 4 * mm), table(['Figure', 'Value'], export.summary, [3, 2])]
    for t in export.tables:
        story.append(Paragraph(escape(t.title), heading_style))
        story.append(table(t.columns, t.rows, t.widths) if t.rows else Paragraph('Nothing to show.', line_style))

    generated = datetime.now().strftime('%Y-%m-%d %H:%M')

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont('Helvetica', 7)
        canvas.setFillColor(colors.HexColor('#94a3b8'))
        canvas.drawRightString(page[0] - 12 * mm, 7 * mm, f'{export.title} · generated {generated} · page {doc.page}')
        canvas.restoreState()

    buffer = io.BytesIO()
    SimpleDocTemplate(buffer, pagesize=page, leftMargin=12 * mm, rightMargin=12 * mm,
                      topMargin=12 * mm, bottomMargin=14 * mm, title=export.title).build(
        story, onFirstPage=footer, onLaterPages=footer)
    return buffer.getvalue()


def export_response(export: Export, fmt: str) -> Response:
    """The export as a download: fmt is 'pdf' or 'excel'."""
    stem = re.sub(r'[^A-Za-z0-9_-]+', '_', export.filename).strip('_') or 'export'
    if fmt == 'pdf':
        content, media, ext = to_pdf(export), 'application/pdf', 'pdf'
    else:
        content = to_excel(export)
        media, ext = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', 'xlsx'
    return Response(content, media_type=media,
                    headers={'Content-Disposition': f'attachment; filename="{stem}.{ext}"'})
